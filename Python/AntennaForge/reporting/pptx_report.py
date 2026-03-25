"""
PowerPoint report generation.
Produces a .pptx slide deck from a directory of pattern CSVs.
"""
from __future__ import annotations

import glob
import logging
import math
import os
import io
from typing import TYPE_CHECKING

import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

if TYPE_CHECKING:
    from types import ModuleType

logger = logging.getLogger(__name__)

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False

from reporting.report import (
    default_report_config, _compute_summary, _find_xpol,
    _get_classification_text, _sanitise_filename
)
from core.io import read_pattern_file, extract_freq_from_filename


def generate_pptx_report(input_dir: str, output_path: str | None = None,
                         title: str | None = None,
                         report_cfg: dict | None = None) -> str | None:
    """Generate a PowerPoint report from a directory of pattern CSVs."""
    if not HAS_PPTX:
        logger.error("python-pptx is required for PPTX reports.")
        return None

    if report_cfg is None:
        report_cfg = default_report_config()
    
    if title is None:
        title = report_cfg.get("cover_title", "").strip()
        if not title:
            title = os.path.basename(os.path.normpath(input_dir))
            
    if output_path is None:
        stem = _sanitise_filename(title)
        output_path = os.path.join(input_dir, f"{stem}.pptx")

    # ── Discover CSV files ────────────────────────────────────────
    csv_files = sorted([
        *glob.glob(os.path.join(input_dir, '*.csv')),
        *glob.glob(os.path.join(input_dir, '*.dat'))
    ])
    csv_files = [f for f in csv_files if '_xpol' not in os.path.basename(f)]
    if not csv_files:
        logger.warning("No pattern CSVs found in %s", input_dir)
        return None

    # ── Build summaries ───────────────────────────────────────────
    summaries: list[dict] = []
    for fp in csv_files:
        try:
            az, el, data = read_pattern_file(fp)
            arr = np.array(data)
            s = _compute_summary(fp, az, el, data, arr, np)
            summaries.append(s)
        except Exception:
            logger.exception("Failed to read %s", fp)
    
    if not summaries:
        return None
    
    summaries.sort(key=lambda s: s['freq'] if s['freq'] is not None else 0)

    # ── Create Presentation ───────────────────────────────────────
    prs = Presentation()
    
    # Helper to add classification footer
    cls_text = _get_classification_text(report_cfg)

    def _add_banners(slide):
        if not cls_text:
            return
        # Header
        txBox = slide.shapes.add_textbox(
            Inches(0), Inches(0), Inches(10), Inches(0.25))
        tf = txBox.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.text = cls_text
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(10)
        p.font.bold = True
        # Footer
        txBox = slide.shapes.add_textbox(
            Inches(0), Inches(7.25), Inches(10), Inches(0.25))
        tf = txBox.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.text = cls_text
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(10)
        p.font.bold = True

    # 1. Title Slide
    slide = prs.slides.add_slide(prs.slide_layouts[0]) # Title Slide
    slide.shapes.title.text = title
    subtitle = report_cfg.get("cover_subtitle", "Antenna Pattern Report")
    slide.placeholders[1].text = f"{subtitle}\n{len(summaries)} Frequency Slices"
    _add_banners(slide)

    # 2. Summary Table (as image for simplicity)
    if report_cfg.get("summary_table", True):
        _add_summary_slide(prs, summaries, _add_banners)

    # 3. Gain vs Freq
    if report_cfg.get("gain_vs_freq", True):
        _add_gain_trend_slide(prs, summaries, _add_banners)

    # 4. Beamwidth vs Freq (Separate Slides)
    if report_cfg.get("bw_vs_freq", True):
        _add_bw_trend_slide(prs, summaries, "Azimuth", _add_banners,
                            report_cfg)
        _add_bw_trend_slide(prs, summaries, "Elevation", _add_banners,
                            report_cfg)

    # 5. Overlay Cuts (4 Separate Slides)
    if report_cfg.get("overlay_cuts", True):
        _add_overlay_slide(prs, summaries, "Azimuth Cartesian", _add_banners)
        _add_overlay_slide(prs, summaries, "Elevation Cartesian", _add_banners)
        _add_overlay_slide(prs, summaries, "Azimuth Polar", _add_banners)
        _add_overlay_slide(prs, summaries, "Elevation Polar", _add_banners)

    # 6. Per Frequency
    for s in summaries:
        freq_str = f"{s['freq']:.2f} MHz" if s['freq'] else os.path.basename(s['file'])
        
        if report_cfg.get("per_freq_heatmap"):
            _add_heatmap_slide(prs, s, freq_str, report_cfg, _add_banners)
        
        if report_cfg.get("per_freq_polar"):
            _add_single_polar_slide(
                prs, s, "Azimuth", freq_str, report_cfg, _add_banners)
            _add_single_polar_slide(
                prs, s, "Elevation", freq_str, report_cfg, _add_banners)
            
        if report_cfg.get("per_freq_3d"):
            _add_3d_slide(prs, s, freq_str, report_cfg, _add_banners)

    prs.save(output_path)
    logger.info("PPTX report saved: %s", output_path)
    return output_path


def _add_plot_to_slide(slide, fig, left, top, width):
    image_stream = io.BytesIO()
    fig.savefig(image_stream, format='png', dpi=150, bbox_inches='tight')
    image_stream.seek(0)
    slide.shapes.add_picture(image_stream, left, top, width=width)
    plt.close(fig)

def _add_summary_slide(prs, summaries, footer_fn):
    # Generate table image
    headers = ['Freq (MHz)', 'Peak (dBi)', 'Bore (dBi)', 'Az BW', 'El BW', 'F/B (dB)']
    rows = []
    for s in summaries:
        rows.append([
            f"{s['freq']:.1f}" if s['freq'] else "?",
            f"{s['peak_gain']:.1f}",
            f"{s['bore_gain']:.1f}",
            str(s['az_bw']),
            str(s['el_bw']),
            f"{s['ftb']:.1f}" if s['ftb'] is not None else "-"
        ])
    
    # Chunking for slides if too many rows
    chunk_size = 20
    for i in range(0, len(rows), chunk_size):
        chunk = rows[i:i+chunk_size]
        
        fig, ax = plt.subplots(figsize=(10, 6))
        ax.axis('off')
        table = ax.table(cellText=chunk, colLabels=headers, loc='center', cellLoc='center')
        table.auto_set_font_size(False)
        table.set_fontsize(10)
        table.scale(1, 1.5)
        
        slide = prs.slides.add_slide(prs.slide_layouts[1])  # Title and Content
        slide.shapes.title.text = "Summary Table"
        _add_plot_to_slide(slide, fig, Inches(0.725), Inches(1.8), Inches(8.55))
        
        if slide.placeholders[1]:
            tf = slide.placeholders[1].text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "This table summarizes the key performance metrics for each frequency slice."
            p.font.size = Pt(12)

        footer_fn(slide)

def _add_gain_trend_slide(prs, summaries, footer_fn):
    freqs = [s['freq'] for s in summaries if s['freq'] is not None]
    if not freqs: return
    peaks = [s['peak_gain'] for s in summaries if s['freq'] is not None]
    
    fig, ax = plt.subplots(figsize=(8, 5))
    ax.plot(freqs, peaks, 'b-o', linewidth=2)
    ax.set_xlabel('Frequency (MHz)')
    ax.set_ylabel('Gain (dBi)')
    ax.grid(True, alpha=0.3)
    ax.set_title('Peak Gain vs Frequency')
    
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Gain vs Frequency"
    _add_plot_to_slide(slide, fig, Inches(1.2), Inches(1.8), Inches(7.6))

    if slide.placeholders[1]:
        tf = slide.placeholders[1].text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.text = "Peak gain (dBi) plotted against frequency (MHz)."
        p.font.size = Pt(12)

    footer_fn(slide)

def _add_bw_trend_slide(prs, summaries, plane, footer_fn,
                        report_cfg=None):
    freqs = [s['freq'] for s in summaries if s['freq'] is not None]
    if not freqs: return

    _is_omni = (report_cfg and
                str(report_cfg.get('_antenna_type', '')).lower() == 'omni')
    _use_plain = _is_omni and plane == "Azimuth"

    key = 'az_bw' if plane == "Azimuth" else 'el_bw'
    bws = [s[key] for s in summaries if s['freq'] is not None]
    color = 'b' if plane == "Azimuth" else 'r'
    _bw_label = "Az Beamwidth" if _use_plain else f'{plane} Beamwidth'

    fig, ax = plt.subplots(figsize=(8, 5))
    ax.plot(freqs, bws, f'{color}-o', linewidth=2)
    ax.set_xlabel('Frequency (MHz)')
    ax.set_ylabel(f'{plane} BW (deg)')
    ax.grid(True, alpha=0.3)
    ax.set_title(f'{_bw_label} vs Frequency')

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = f"{_bw_label} Trend"
    _add_plot_to_slide(slide, fig, Inches(1.2), Inches(1.8), Inches(7.6))

    if slide.placeholders[1]:
        tf = slide.placeholders[1].text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.text = f"{_bw_label} (degrees) plotted against frequency (MHz)."
        p.font.size = Pt(12)

    footer_fn(slide)

def _add_overlay_slide(prs, summaries, mode, footer_fn):
    fig = plt.figure(figsize=(8, 5))
    ax = fig.add_subplot(111, projection='polar' if 'Polar' in mode else None)
    
    colors = plt.cm.turbo(np.linspace(0, 1, max(len(summaries), 1)))
    
    for i, s in enumerate(summaries):
        arr = s['arr']
        if 'Azimuth' in mode:
            cut = arr[s['el0'], :]
            angles = s['az']
        else:
            cut = arr[:, s['az0']]
            angles = s['el']
            
        lbl = f"{s['freq']:.0f}" if s['freq'] else ""
        
        if 'Polar' in mode:
            peak = float(arr.max())
            floor = peak - 40
            norm = np.clip(cut, floor, peak)
            norm = (norm - floor) / max(peak - floor, 1)
            ax.plot(np.radians(angles), norm, color=colors[i], linewidth=1, label=lbl)
        else:
            ax.plot(angles, cut, color=colors[i], linewidth=1, label=lbl)
            
    if 'Polar' in mode:
        ax.set_theta_zero_location('N')
        ax.set_theta_direction(-1)
    else:
        ax.set_xlabel('Angle (deg)')
        ax.set_ylabel('Gain (dBi)')
        ax.grid(True, alpha=0.3)
        
    # Legend only if few curves
    if len(summaries) <= 10:
        ax.legend(fontsize=8)
        
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = f"Overlay: {mode}"
    
    width = Inches(7.6)
    left = Inches(1.2)
    if 'Polar' in mode:
        width = Inches(6.8)
        left = Inches(1.6)
    _add_plot_to_slide(slide, fig, left, Inches(1.8), width)

    if slide.placeholders[1]:
        tf = slide.placeholders[1].text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.text = f"Overlay of {mode} cuts for all frequencies."
        p.font.size = Pt(12)

    footer_fn(slide)

def _add_heatmap_slide(prs, s, freq_str, report_cfg, footer_fn):
    az = s['az']
    el = s['el']
    arr = s['arr']
    dyn_range = report_cfg.get("dynamic_range_db", 60)
    
    fig = plt.figure(figsize=(10, 5.5))
    gs = fig.add_gridspec(1, 2, width_ratios=[1.5, 1])
    
    # Heatmap
    ax1 = fig.add_subplot(gs[0])
    vmin = max(arr.min(), arr.max() - dyn_range)
    AZ, EL = np.meshgrid(az, el)
    im = ax1.pcolormesh(AZ, EL, arr, shading='gouraud', cmap='jet', vmin=vmin, vmax=arr.max())
    ax1.set_xlabel('Az (deg)')
    ax1.set_ylabel('El (deg)')
    ax1.set_title('Heatmap')
    fig.colorbar(im, ax=ax1, shrink=0.7, label='dBi')
    
    # Cuts
    ax2 = fig.add_subplot(gs[1])
    az_cut = arr[s['el0'], :]
    el_cut = arr[:, s['az0']]
    ax2.plot(az, az_cut, 'b-', label='Az Cut')
    ax2.plot(el, el_cut, 'r-', label='El Cut')
    ax2.set_xlabel('Angle (deg)')
    ax2.set_ylabel('Gain (dBi)')
    ax2.grid(True, alpha=0.3)
    ax2.legend()
    ax2.set_title('Principal Cuts')
    
    plt.tight_layout()
    
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = f"Pattern @ {freq_str}"
    _add_plot_to_slide(slide, fig, Inches(0.725), Inches(1.8), Inches(8.55))
    
    if slide.placeholders[1]:
        tf = slide.placeholders[1].text_frame
        tf.clear()
        p1 = tf.paragraphs[0]
        p1.text = "2D heatmap of the full antenna pattern with principal cuts."
        p1.font.size = Pt(12)
        
        p2 = tf.add_paragraph()
        stats = (f"Peak: {s['peak_gain']:.1f} dBi   |   "
                 f"Az BW: {s['az_bw']}°   |   "
                 f"El BW: {s['el_bw']}°")
        p2.text = stats
        p2.font.size = Pt(11)
        p2.font.bold = True

    footer_fn(slide)

def _add_single_polar_slide(prs, s, plane, freq_str, report_cfg, footer_fn):
    """Add a slide with a single polar plot for the given plane."""
    az = s['az']
    el = s['el']
    arr = s['arr']
    polar_range = report_cfg.get("polar_range_db", 40)
    
    fig = plt.figure(figsize=(6, 4.5))
    ax = fig.add_subplot(111, projection='polar')
    
    peak = float(arr.max())
    floor = peak - polar_range
    
    if plane == "Azimuth":
        cut = arr[s['el0'], :]
        angles = az
        color = 'b'
    else: # Elevation
        cut = arr[:, s['az0']]
        angles = el
        color = 'r'

    plot_data = np.clip(cut, floor, peak) - floor
    ax.plot(np.radians(angles), plot_data, color=color)
    ax.set_title(plane)
    ax.set_theta_zero_location('N')
    ax.set_theta_direction(-1)
    ax.set_yticklabels([])

    plt.tight_layout()
    
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = f"{plane} Polar Plot @ {freq_str}"
    _add_plot_to_slide(slide, fig, Inches(2.45), Inches(1.8), Inches(5.1))

    if slide.placeholders[1]:
        tf = slide.placeholders[1].text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.text = f"Polar plot of the {plane} cut. The radial axis shows gain in dB, normalized to the peak."
        p.font.size = Pt(12)

    footer_fn(slide)

def _add_3d_slide(prs, s, freq_str, report_cfg, footer_fn):
    az = s['az']
    el = s['el']
    arr = s['arr']
    dyn_range = report_cfg.get("dynamic_range_db", 60)
    
    step = max(1, len(az) // 60)
    az_sub = az[::step]
    el_sub = el[::step]
    arr_sub = arr[::step, ::step]
    AZ, EL = np.meshgrid(az_sub, el_sub)
    vmin = max(arr_sub.min(), arr_sub.max() - dyn_range)
    arr_c = np.clip(arr_sub, vmin, arr_sub.max())
    
    fig = plt.figure(figsize=(8, 5))
    ax = fig.add_subplot(111, projection='3d')
    ax.plot_surface(AZ, EL, arr_c, cmap='jet', rstride=1, cstride=1)
    ax.set_xlabel('Az')
    ax.set_ylabel('El')
    ax.set_zlabel('Gain')
    
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = f"3D Surface @ {freq_str}"
    _add_plot_to_slide(slide, fig, Inches(1.2), Inches(1.8), Inches(7.6))

    if slide.placeholders[1]:
        tf = slide.placeholders[1].text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.text = "3D surface plot of the antenna pattern, showing gain over azimuth and elevation."
        p.font.size = Pt(12)

    footer_fn(slide)
